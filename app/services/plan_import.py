"""Nhập file Excel kế hoạch SX (.xlsb/.xlsx) — cấu trúc theo file "XN CHỐT".

Sheet sử dụng:
  KẾ HOẠCH     : PO đã xếp chuyền/thời gian (planning_status PLANNED)
  PO MỚI       : PO chưa lên KH (planning_status UNPLANNED) — có thể đã biết XN hoặc chưa
  PO MAY XONG  : đã may xong, chờ xuất      (chỉ đếm)
  ĐÃ XUẤT      : lịch sử giao hàng          (chỉ đếm)
  LAO ĐỘNG     : công nhân có mặt theo Xí nghiệp/Tổ (Nhân sự)
"""

import logging
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterator

from sqlalchemy.orm import Session

from app.models.core import Factory
from app.models.data import LaborHeadcount, PlanImportBatch, PlanRow, SyncRun
from app.services.rules import assess_factory, classify_row, excel_serial_to_date, strip_accents, to_float
from app.services.sync_core import add_items, fail_run, finish_run, start_run

log = logging.getLogger(__name__)

# Vị trí cột (0-based) trong các sheet KẾ HOẠCH / PO MỚI / PO MAY XONG / ĐÃ XUẤT
C_FAC, C_LINE, C_PO_DATE, C_SPORT, C_SEASON, C_PO, C_STYLE, C_MODEL, C_DESC, C_CUSTOMER = range(10)
C_QTY, C_WORKER, C_CAPACITY, C_TOTAL_DAY = 10, 11, 12, 13
C_FABRIC, C_ACC = 15, 16
C_BEGIN, C_END = 22, 23
C_WH = 26
C_CHD, C_EHD = 28, 29
C_GAP, C_STATUS, C_NOTE, C_DEST = 31, 32, 33, 34
C_PROCESS = 37


def _g(row: list, idx: int):
    return row[idx] if idx < len(row) else None


def _s(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _norm_sheet(name: str) -> str:
    return strip_accents(name).upper().strip()


class Workbook:
    def __init__(self, path: Path):
        self.path = path
        self.ext = path.suffix.lower()
        self._ctx = None
        self._wb = None

    def __enter__(self):
        if self.ext == ".xlsb":
            from pyxlsb import open_workbook

            self._ctx = open_workbook(str(self.path))
            self._wb = self._ctx.__enter__()
            self._names = list(self._wb.sheets)
        elif self.ext in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook

            self._wb = load_workbook(str(self.path), read_only=True, data_only=True)
            self._names = list(self._wb.sheetnames)
        else:
            raise ValueError("Chỉ hỗ trợ file .xlsb, .xlsx, .xlsm")
        return self

    def __exit__(self, *exc):
        if self._ctx is not None:
            self._ctx.__exit__(*exc)
        elif self._wb is not None:
            self._wb.close()

    def find(self, wanted: str) -> str | None:
        target = _norm_sheet(wanted)
        for n in self._names:
            if _norm_sheet(n) == target:
                return n
        return None

    def rows(self, name: str) -> Iterator[list]:
        if self.ext == ".xlsb":
            with self._wb.get_sheet(name) as sheet:
                for row in sheet.rows():
                    yield [c.v for c in row]
        else:
            for row in self._wb[name].iter_rows(values_only=True):
                yield list(row)


@dataclass
class ParsedPlan:
    planned: list[dict] = field(default_factory=list)
    new: list[dict] = field(default_factory=list)
    sewn_count: int = 0
    sewn_qty: float = 0
    shipped_count: int = 0
    shipped_qty: float = 0
    labor: list[dict] = field(default_factory=list)
    labor_as_of: str = ""
    sheets_found: list[str] = field(default_factory=list)


def _parse_plan_rows(rows: Iterator[list], planning_status: str) -> list[dict]:
    out = []
    for i, row in enumerate(rows):
        if i < 2:  # dòng tiêu đề + dòng đánh số cột
            continue
        qty = to_float(_g(row, C_QTY))
        if not qty or qty <= 0:
            continue
        fabric = _g(row, C_FABRIC)
        fabric_text = fabric if isinstance(fabric, str) else None
        chd = excel_serial_to_date(_g(row, C_CHD))
        gap = to_float(_g(row, C_GAP))
        # Dòng dự báo/đào tạo không có CHD nên EHD/CHD vô nghĩa (ra hàng chục nghìn ngày) -> bỏ qua
        if chd is None or gap is None or abs(gap) > 3650:
            gap = None
        status_text = _s(_g(row, C_STATUS))
        note = _s(_g(row, C_NOTE))
        risk, reason = classify_row(gap, status_text, note, fabric_text)
        fac_num = to_float(_g(row, C_FAC))
        out.append(
            dict(
                planning_status=planning_status,
                fac_num=int(fac_num) if fac_num is not None and fac_num == int(fac_num) else None,
                fac_raw=_s(_g(row, C_FAC))[:20],
                line_raw=_s(_g(row, C_LINE))[:60],
                po_date=excel_serial_to_date(_g(row, C_PO_DATE)),
                sport=_s(_g(row, C_SPORT))[:60],
                season=_s(_g(row, C_SEASON))[:30],
                po_number=_s(_g(row, C_PO))[:60],
                style_cc=_s(_g(row, C_STYLE))[:60],
                model_code=_s(_g(row, C_MODEL))[:60],
                description=_s(_g(row, C_DESC))[:300],
                customer=_s(_g(row, C_CUSTOMER))[:100],
                quantity=qty,
                worker=to_float(_g(row, C_WORKER)),
                capacity=to_float(_g(row, C_CAPACITY)),
                total_day=to_float(_g(row, C_TOTAL_DAY)),
                begin_prod_date=excel_serial_to_date(_g(row, C_BEGIN)),
                end_prod_date=excel_serial_to_date(_g(row, C_END)),
                warehouse_date=excel_serial_to_date(_g(row, C_WH)),
                chd=chd,
                ehd_etd=excel_serial_to_date(_g(row, C_EHD)),
                gap_days=gap,
                status_text=status_text[:40],
                note=(note or (fabric_text or ""))[:400],
                destination=_s(_g(row, C_DEST))[:100],
                process=_s(_g(row, C_PROCESS))[:60],
                risk=risk,
                risk_reason=reason[:200],
            )
        )
    return out


def _count_rows(rows: Iterator[list]) -> tuple[int, float]:
    count, total = 0, 0.0
    for i, row in enumerate(rows):
        if i < 2:
            continue
        qty = to_float(_g(row, C_QTY))
        if qty and qty > 0:
            count += 1
            total += qty
    return count, total


def parse_plan_workbook(path: Path) -> ParsedPlan:
    parsed = ParsedPlan()
    with Workbook(path) as wb:
        plan_sheet = wb.find("KẾ HOẠCH")
        if plan_sheet is None:
            raise ValueError("Không tìm thấy sheet 'KẾ HOẠCH' trong file")
        parsed.sheets_found.append(plan_sheet)
        parsed.planned = _parse_plan_rows(wb.rows(plan_sheet), "PLANNED")

        if (name := wb.find("PO MỚI")) is not None:
            parsed.sheets_found.append(name)
            parsed.new = _parse_plan_rows(wb.rows(name), "UNPLANNED")
        if (name := wb.find("PO MAY XONG")) is not None:
            parsed.sheets_found.append(name)
            parsed.sewn_count, parsed.sewn_qty = _count_rows(wb.rows(name))
        if (name := wb.find("ĐÃ XUẤT")) is not None:
            parsed.sheets_found.append(name)
            parsed.shipped_count, parsed.shipped_qty = _count_rows(wb.rows(name))
        if (name := wb.find("LAO ĐỘNG")) is not None:
            parsed.sheets_found.append(name)
            for i, row in enumerate(wb.rows(name)):
                if i == 0:
                    parsed.labor_as_of = _s(_g(row, 2))
                    continue
                xn, team, count = to_float(_g(row, 0)), _s(_g(row, 1)), to_float(_g(row, 2))
                if xn is None or count is None:
                    continue
                parsed.labor.append(dict(fac_num=int(xn), team=team[:20], headcount=int(count)))
    return parsed


def _sync(db: Session, run: SyncRun, path: Path, filename: str, username: str) -> None:
    parsed = parse_plan_workbook(path)
    fmap = {f.sql_xn_id: f.id for f in db.query(Factory).filter(Factory.sql_xn_id.isnot(None)).all()}

    db.query(PlanImportBatch).update({PlanImportBatch.is_current: False})
    db.query(PlanRow).delete()
    db.query(LaborHeadcount).delete()

    batch = PlanImportBatch(
        filename=filename,
        imported_by=username,
        is_current=True,
        sync_run_id=run.id,
        summary={
            "planned": len(parsed.planned),
            "new": len(parsed.new),
            "unplanned_known": 0,
            "unplanned_unassigned": 0,
            "sewn_count": parsed.sewn_count,
            "sewn_qty": parsed.sewn_qty,
            "shipped_count": parsed.shipped_count,
            "shipped_qty": parsed.shipped_qty,
            "labor_as_of": parsed.labor_as_of,
            "sheets": parsed.sheets_found,
        },
    )
    db.add(batch)
    db.flush()

    # Ba chiều độc lập: planning_status (theo sheet) · factory_assignment (đã biết XN?) · mapping_status (dữ liệu nguồn có vấn đề?)
    warnings, matched = [], 0
    stats = {
        "PLANNED": {"matched": 0, "unmatched": 0},
        "UNPLANNED": {"matched": 0, "unmatched": 0},
    }
    unplanned_known = unplanned_unassigned = 0
    rows = []
    for r in parsed.planned + parsed.new:
        fid = fmap.get(r["fac_num"]) if r["fac_num"] is not None else None
        assignment, mapping, note = assess_factory(r["planning_status"], fid is not None, r["fac_raw"])
        if mapping == "WARNING":
            stats[r["planning_status"]]["unmatched"] += 1
            warnings.append(
                (
                    f"PO {r['po_number'] or '?'} / style {r['style_cc']}",
                    f"{'KẾ HOẠCH' if r['planning_status'] == 'PLANNED' else 'PO MỚI'}: {note} — {r['customer']} · SL {r['quantity']:,.0f} · {r['description'][:60]}",
                    {"po": r["po_number"], "qty": r["quantity"], "fac": r["fac_raw"]},
                )
            )
        else:
            stats[r["planning_status"]]["matched"] += 1
            matched += 1
        if r["planning_status"] == "UNPLANNED":
            unplanned_known += assignment == "KNOWN"
            unplanned_unassigned += assignment == "UNASSIGNED"
        row = {k: v for k, v in r.items() if k != "fac_num"}
        row.update(factory_id=fid, batch_id=batch.id, factory_assignment=assignment, mapping_status=mapping, mapping_note=note)
        rows.append(row)
    unmatched_plan = warnings
    for i in range(0, len(rows), 2000):
        db.bulk_insert_mappings(PlanRow, rows[i : i + 2000])

    labor_rows, unmatched_labor, labor_matched = [], [], 0
    for r in parsed.labor:
        fid = fmap.get(r["fac_num"])
        if fid is None:
            unmatched_labor.append((f"XN={r['fac_num']} Tổ={r['team']}", "Lao động thuộc Xí nghiệp không có trong danh mục", dict(r)))
        else:
            labor_matched += 1
        labor_rows.append(dict(batch_id=batch.id, factory_id=fid, team=r["team"], headcount=r["headcount"], as_of_text=parsed.labor_as_of))
    if labor_rows:
        db.bulk_insert_mappings(LaborHeadcount, labor_rows)

    batch.summary = {**batch.summary, "unplanned_known": unplanned_known, "unplanned_unassigned": unplanned_unassigned}
    add_items(db, run.id, "UNMATCHED", "KẾ HOẠCH / PO MỚI", unmatched_plan)
    add_items(db, run.id, "UNMATCHED", "LAO ĐỘNG", unmatched_labor)

    run.total_records = len(rows) + len(labor_rows)
    run.matched = matched + labor_matched
    run.unmatched = len(unmatched_plan) + len(unmatched_labor)
    run.ambiguous = 0
    run.updated_rows = len(rows) + len(labor_rows)
    run.summary = {
        "filename": filename,
        "objects": [
            {"name": "KẾ HOẠCH", "read": len(parsed.planned), **stats["PLANNED"]},
            {"name": "PO MỚI", "read": len(parsed.new), **stats["UNPLANNED"]},
            {"name": "PO MAY XONG", "read": parsed.sewn_count, "matched": parsed.sewn_count, "unmatched": 0},
            {"name": "ĐÃ XUẤT", "read": parsed.shipped_count, "matched": parsed.shipped_count, "unmatched": 0},
            {"name": "LAO ĐỘNG", "read": len(parsed.labor), "matched": labor_matched, "unmatched": len(unmatched_labor)},
        ],
    }
    db.commit()


def run_plan_import(db: Session, path: Path, filename: str, username: str) -> SyncRun:
    run = start_run(db, "PLAN_EXCEL", "MANUAL", username)
    t0 = time.perf_counter()
    try:
        _sync(db, run, path, filename, username)
        return finish_run(db, run, t0)
    except Exception as exc:  # noqa: BLE001
        return fail_run(db, run.id, t0, exc)
